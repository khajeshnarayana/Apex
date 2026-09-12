"""
inspect_resumes.py — quick look at what's actually inside the dataset.

Person 2 utility: NOT the real parser (that's Person 1's job).
This just dumps text so we can see real phrasing, real section headers,
and build realistic test cases against actual resume content.

Usage:
    pip install python-docx
    python inspect_resumes.py

Edit RESUME_DIR and FILES_TO_INSPECT below if your paths differ.
"""

import os
from docx import Document

RESUME_DIR = r"D:\DEV\Apex\data\Dummy Resumes"

# Deliberately varied domains so we can test ranking spread.
# Change these names if they don't match what's in your folder —
# run with LIST_ALL = True first to see available files.
FILES_TO_INSPECT = [
    "SDE_Resume_1_Aditya_Joshi.docx",
    "Python_Developer_Resume_1_Karan_Malhotra.docx",
    "App_Developer_Resume_1_Siddharth_Rao.docx",
    "Sales_Resume_1_Sahil_Khanna.docx",
    "Video_Editing_Resume_1_Rahul_Saxena.docx",
]

LIST_ALL = False       # set True to just list every .docx filename and exit
MAX_CHARS = 2500       # truncate long resumes so output stays readable


def list_all_docx():
    files = sorted(f for f in os.listdir(RESUME_DIR) if f.lower().endswith(".docx"))
    print(f"Found {len(files)} .docx files:\n")
    for f in files:
        print("  ", f)


def read_docx(path):
    """Pull all non-empty paragraph text out of a .docx file."""
    doc = Document(path)
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Tables often hold skills in resume templates — grab those too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)

    return lines


def guess_headers(lines):
    """
    Rough heuristic for spotting section headers — useful intel for Person 1.
    A header is usually short, and often all-caps or title-case.
    """
    headers = []
    for line in lines:
        words = line.split()
        if 1 <= len(words) <= 4 and len(line) < 40:
            if line.isupper() or line.istitle():
                headers.append(line)
    return headers


def inspect(filename):
    path = os.path.join(RESUME_DIR, filename)

    if not os.path.exists(path):
        print(f"\n!! NOT FOUND: {filename}")
        print("   (run with LIST_ALL = True to see actual filenames)")
        return

    print("\n" + "=" * 70)
    print(f"FILE: {filename}")
    print("=" * 70)

    lines = read_docx(path)
    full_text = "\n".join(lines)

    print(f"\n--- DETECTED HEADER CANDIDATES (intel for Person 1) ---")
    headers = guess_headers(lines)
    for h in headers:
        print("   *", h)

    print(f"\n--- TEXT ({len(full_text)} chars total) ---\n")
    print(full_text[:MAX_CHARS])
    if len(full_text) > MAX_CHARS:
        print(f"\n... [truncated, {len(full_text) - MAX_CHARS} more chars]")


if __name__ == "__main__":
    if not os.path.isdir(RESUME_DIR):
        print(f"Directory not found: {RESUME_DIR}")
        print("Edit RESUME_DIR at the top of this file.")
        raise SystemExit(1)

    if LIST_ALL:
        list_all_docx()
    else:
        for filename in FILES_TO_INSPECT:
            inspect(filename)